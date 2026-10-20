#!/usr/bin/env python3
"""
Captação de lojas (Brasil): roda o scrape.py para vários nichos em várias cidades e junta tudo
num único relatório de contatos, sem duplicatas.

Colunas do relatório: nome, categoria, telefone, whatsapp, link_whatsapp, email, site, endereco,
cidade, uf, nota, avaliacoes, busca.

Uso:
    python3 scripts/captar_lojas.py --cidade "Uberlandia, MG"
    python3 scripts/captar_lojas.py --cidade "Uberlandia, MG" --cidade "Goiania, GO" --depth 3
    python3 scripts/captar_lojas.py --cidades-file examples/cidades.example.txt
    python3 scripts/captar_lojas.py --cidade "Uberlandia, MG;-18.9186;-48.2772"   # sem geocoding
    python3 scripts/captar_lojas.py --cidade "Uberlandia, MG" --nichos "ferragens,parafusos"

Saída: relatorio-lojas-<data>.xlsx se o openpyxl estiver instalado (pip install openpyxl),
senão relatorio-lojas-<data>.csv separado por ";" (abre direto no Excel em português).

Observação: o Google Maps mostra o nome fantasia, não a razão social. A razão social sai da
consulta do CNPJ, numa etapa separada.
"""
import argparse, csv, datetime, os, re, subprocess, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))

NICHOS = [
    "loja de material de construção",
    "loja de material elétrico",
    "loja de eletrodomésticos",
    "loja de ferramentas",
    "loja de ferragens",
    "loja de parafusos",
    "loja de utilidades domésticas",
    "loja de eletrônicos",
    "loja de informática",
]

COLUNAS = ["nome", "categoria", "telefone", "whatsapp", "link_whatsapp", "email", "site",
           "endereco", "cidade", "uf", "nota", "avaliacoes", "busca"]

# "..., Uberlândia - MG, 38400-000" ou "..., Uberlândia - MG"
RE_CIDADE_UF = re.compile(r"([^,]+?)\s+-\s+([A-Z]{2})(?:\s*,\s*\d{5}-?\d{3})?(?:\s*,\s*Brazil|\s*,\s*Brasil)?\s*$")


def digitos(tel):
    d = re.sub(r"\D", "", tel or "")
    if d.startswith("55") and len(d) in (12, 13):
        d = d[2:]
    return d.lstrip("0")


def whatsapp(tel):
    """Celular BR (DDD + 9 + 8 dígitos) vira número de WhatsApp provável. Fixo não."""
    d = digitos(tel)
    if len(d) == 11 and d[2] == "9":
        return f"+55 ({d[:2]}) {d[2:7]}-{d[7:]}", f"https://wa.me/55{d}"
    return "", ""


def cidade_uf(endereco, cidade_padrao):
    m = RE_CIDADE_UF.search((endereco or "").strip())
    if m:
        cidade = m.group(1).split(" - ")[-1].strip()
        return cidade, m.group(2)
    nome, _, uf = cidade_padrao.partition(",")
    return nome.strip(), uf.strip().upper()


def parse_cidade(linha):
    """'Uberlandia, MG' ou 'Uberlandia, MG;-18.91;-48.27' -> (nome, lat, lon)."""
    partes = [p.strip() for p in linha.split(";")]
    if len(partes) == 3:
        return partes[0], partes[1], partes[2]
    return partes[0], None, None


def rodar_cidade(cidade, lat, lon, nichos, depth, sem_email, pasta):
    nome_cidade = cidade.split(",")[0].strip()
    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False, encoding="utf-8") as f:
        for n in nichos:
            f.write(f"{n} em {cidade}\n")
        kw_file = f.name
    slug = re.sub(r"\W+", "_", nome_cidade.lower()).strip("_")
    out = os.path.join(pasta, f"bruto-{slug}.csv")
    cmd = [sys.executable, os.path.join(HERE, "scrape.py"), "--keywords-file", kw_file,
           "--depth", str(depth), "--out", out, "--max-time", str(max(600, 120 * len(nichos)))]
    cmd += [lat, lon] if lat and lon else ["--city", f"{cidade}, Brasil"]
    if sem_email:
        cmd.append("--no-email")
    print(f"\n=== {cidade}: {len(nichos)} nichos ===")
    try:
        r = subprocess.run(cmd)
    finally:
        os.unlink(kw_file)
    if r.returncode != 0 or not os.path.exists(out):
        print(f"✗ {cidade}: falhou (veja as mensagens acima).")
        return []
    with open(out, newline="", encoding="utf-8") as f:
        return [dict(row, _cidade=cidade) for row in csv.DictReader(f)]


def montar_relatorio(linhas):
    vistos, saida = set(), []
    for r in linhas:
        tel = r.get("phone", "")
        chave = digitos(tel) or (r.get("title", "").lower(), r.get("address", "").lower())
        if chave in vistos:
            continue
        vistos.add(chave)
        wa, link = whatsapp(tel)
        cid, uf = cidade_uf(r.get("address", ""), r["_cidade"])
        saida.append({
            "nome": r.get("title", ""), "categoria": r.get("category", ""), "telefone": tel,
            "whatsapp": wa, "link_whatsapp": link, "email": r.get("emails", ""),
            "site": r.get("website", ""), "endereco": r.get("address", ""), "cidade": cid,
            "uf": uf, "nota": r.get("review_rating", ""), "avaliacoes": r.get("review_count", ""),
            "busca": r["_cidade"],
        })
    saida.sort(key=lambda x: (x["uf"], x["cidade"], x["categoria"], x["nome"]))
    return saida


def salvar(linhas, base):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        out = base + ".csv"
        with open(out, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=COLUNAS, delimiter=";")
            w.writeheader()
            w.writerows(linhas)
        return out
    wb = Workbook()
    ws = wb.active
    ws.title = "Contatos"
    ws.append(COLUNAS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in linhas:
        ws.append([r[c] for c in COLUNAS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, larg in zip("ABCDEFGHIJKLM", (34, 24, 18, 20, 30, 30, 30, 50, 20, 5, 6, 10, 22)):
        ws.column_dimensions[col].width = larg
    resumo = wb.create_sheet("Resumo")
    resumo.append(["cidade", "uf", "lojas", "com whatsapp", "com email"])
    grupos = {}
    for r in linhas:
        g = grupos.setdefault((r["cidade"], r["uf"]), [0, 0, 0])
        g[0] += 1
        g[1] += bool(r["whatsapp"])
        g[2] += bool(r["email"])
    for (cid, uf), (n, wa, em) in sorted(grupos.items(), key=lambda x: -x[1][0]):
        resumo.append([cid, uf, n, wa, em])
    out = base + ".xlsx"
    wb.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description="Captação de lojas por nicho e cidade (Google Maps).")
    ap.add_argument("--cidade", action="append", default=[],
                    help='"Cidade, UF" ou "Cidade, UF;lat;lon" (repita para várias)')
    ap.add_argument("--cidades-file", help="arquivo com uma cidade por linha")
    ap.add_argument("--nichos", help="lista separada por vírgula (padrão: todos os 9 nichos)")
    ap.add_argument("--depth", type=int, default=3, help="profundidade da rolagem (padrão 3)")
    ap.add_argument("--sem-email", action="store_true", help="mais rápido: não busca emails")
    ap.add_argument("--saida", default=None, help="nome base do relatório (sem extensão)")
    a = ap.parse_args()

    cidades = list(a.cidade)
    if a.cidades_file:
        with open(a.cidades_file, encoding="utf-8") as f:
            cidades += [l.strip() for l in f if l.strip() and not l.startswith("#")]
    if not cidades:
        ap.error("informe --cidade ou --cidades-file")
    nichos = [n.strip() for n in a.nichos.split(",")] if a.nichos else NICHOS

    if len(cidades) * len(nichos) >= 20 or a.depth >= 10:
        print("⚠️  Busca grande: pode levar bastante tempo e o Google pode limitar seu IP por um tempo.")

    hoje = datetime.date.today().isoformat()
    pasta = os.path.join(os.getcwd(), f"captacao-{hoje}")
    os.makedirs(pasta, exist_ok=True)
    brutos = []
    for c in cidades:
        nome, lat, lon = parse_cidade(c)
        brutos += rodar_cidade(nome, lat, lon, nichos, a.depth, a.sem_email, pasta)

    linhas = montar_relatorio(brutos)
    out = salvar(linhas, a.saida or os.path.join(pasta, f"relatorio-lojas-{hoje}"))
    wa = sum(1 for r in linhas if r["whatsapp"])
    em = sum(1 for r in linhas if r["email"])
    print(f"\n✓ {len(linhas)} lojas únicas | {wa} com WhatsApp | {em} com email")
    print(f"  relatório → {out}")


if __name__ == "__main__":
    main()
